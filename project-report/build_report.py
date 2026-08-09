from __future__ import annotations

import json
import os
import re
import shutil
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "project-report"
SOURCE_PATH = REPORT_DIR / "fastclaw_project_report.md"
TEMPLATE_PATH = Path(
    os.environ.get(
        "FASTCLAW_REPORT_TEMPLATE",
        "/Users/wangzheyu/Downloads/ieee_financial_analysis_assistant_verified.docx",
    )
)
OUTPUT_PATH = REPORT_DIR / "FastClaw_Project_Report_Zheyu_Wang.docx"
ASSET_DIR = REPORT_DIR / "assets"
PAGE_MAP_PATH = REPORT_DIR / "page_map.json"

NAVY = "#17324D"
BLUE = "#1F5A7A"
TEAL = "#2F7D7A"
LIGHT_BLUE = "#EAF3F8"
LIGHT_TEAL = "#EAF6F3"
LIGHT_GRAY = "#F2F4F6"
MID_GRAY = "#D8DEE4"
DARK = "#1F252B"
MUTED = "#59636E"
WHITE = "#FFFFFF"
RED = "#A23B3B"
GREEN = "#2C6E49"
BODY_WIDTH_INCHES = 6.80


def _font_path(bold: bool = False, mono: bool = False) -> str:
    if mono:
        candidates = [
            "/System/Library/Fonts/Supplemental/Courier New Bold.ttf"
            if bold
            else "/System/Library/Fonts/Supplemental/Courier New.ttf",
            "/Library/Fonts/Courier New Bold.ttf" if bold else "/Library/Fonts/Courier New.ttf",
        ]
    else:
        candidates = [
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
            if bold
            else "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError("A compatible macOS TrueType font was not found")


def _font(size: int, bold: bool = False, mono: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(_font_path(bold=bold, mono=mono), size=size)


def _wrap(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def _draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    font: ImageFont.FreeTypeFont,
    fill: str = DARK,
    padding: int = 24,
    line_gap: int = 8,
) -> None:
    left, top, right, bottom = box
    lines = _wrap(draw, text, font, right - left - 2 * padding)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total_height = sum(heights) + line_gap * max(0, len(lines) - 1)
    cursor_y = top + (bottom - top - total_height) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        cursor_x = left + (right - left - (bbox[2] - bbox[0])) / 2
        draw.text((cursor_x, cursor_y), line, font=font, fill=fill)
        cursor_y += height + line_gap


def _rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    body: str,
    *,
    fill: str,
    outline: str = BLUE,
) -> None:
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=4)
    left, top, right, _ = box
    draw.text((left + 24, top + 17), title, font=_font(31, bold=True), fill=NAVY)
    body_box = (left + 12, top + 56, right - 12, box[3] - 8)
    _draw_centered_text(draw, body_box, body, font=_font(24), fill=DARK, padding=18, line_gap=6)


def _arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    color: str = BLUE,
    width: int = 6,
) -> None:
    draw.line((start, end), fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - direction * 18, y2 - 12), (x2 - direction * 18, y2 + 12)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 12, y2 - direction * 18), (x2 + 12, y2 - direction * 18)]
    draw.polygon(points, fill=color)


def create_architecture_figure(path: Path) -> None:
    image = Image.new("RGB", (2400, 1580), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((85, 58), "FastClaw Layered Runtime Architecture", font=_font(48, bold=True), fill=NAVY)
    draw.text(
        (85, 122),
        "Authenticated ingress, bounded orchestration, policy-controlled execution, and tenant-scoped state",
        font=_font(26),
        fill=MUTED,
    )

    layers = [
        (
            "1  Ingress",
            "Web UI  ·  OpenAI-compatible API  ·  WebSocket  ·  Channels  ·  Webhook  ·  Cron",
            LIGHT_BLUE,
        ),
        (
            "2  Identity & Scope",
            "Cookie session / Bearer API key  ·  User identity  ·  Agent ACL  ·  Fail-closed ownership",
            LIGHT_TEAL,
        ),
        (
            "3  Gateway Orchestration",
            "MessageBus  ·  Per-chat FIFO TaskQueue  ·  Lazy UserSpace  ·  Global and internal bounds",
            LIGHT_BLUE,
        ),
        (
            "4  Agent Runtime",
            "Context builder  ·  Streaming ReAct loop  ·  Session repair  ·  Event emitter  ·  Usage collector",
            LIGHT_TEAL,
        ),
        (
            "5  Capability Layer",
            "Built-in tools  ·  MCP servers  ·  JSON-RPC plugins  ·  Skills  ·  Docker/E2B sandbox",
            LIGHT_BLUE,
        ),
        (
            "6  Persistence & Telemetry",
            "SQLite/PostgreSQL  ·  Workspace objects  ·  Sessions  ·  Agent files  ·  Traces and evaluation",
            LIGHT_TEAL,
        ),
    ]
    left, right = 105, 2295
    top = 205
    height = 170
    gap = 48
    for index, (title, body, fill) in enumerate(layers):
        y1 = top + index * (height + gap)
        y2 = y1 + height
        _rounded_box(draw, (left, y1, right, y2), title, body, fill=fill)
        if index < len(layers) - 1:
            _arrow(draw, ((left + right) // 2, y2 + 7), ((left + right) // 2, y2 + gap - 7))

    image.save(path, dpi=(240, 240))


def create_sequence_figure(path: Path) -> None:
    image = Image.new("RGB", (2400, 1500), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((75, 50), "One Multi-Agent Turn: Request, Delegation, and Streaming", font=_font(46, bold=True), fill=NAVY)
    participants = ["Client", "Gateway", "TaskQueue", "Coordinator", "Provider", "MessageBus", "Specialists (3)"]
    x_positions = [130, 480, 830, 1180, 1530, 1880, 2250]
    top = 160
    bottom = 1400
    for label, x in zip(participants, x_positions):
        draw.rounded_rectangle((x - 112, top, x + 112, top + 78), radius=18, fill=LIGHT_BLUE, outline=BLUE, width=3)
        _draw_centered_text(draw, (x - 112, top, x + 112, top + 78), label, font=_font(22, bold=True), padding=8)
        draw.line((x, top + 88, x, bottom), fill=MID_GRAY, width=3)

    events = [
        (0, 1, 300, "authenticated request"),
        (1, 2, 405, "enqueue by chat key"),
        (2, 3, 510, "start root execution"),
        (3, 4, 615, "ChatStream(messages, tools)"),
        (4, 3, 720, "tool call: spawn_subagent"),
        (3, 5, 810, "correlated request + call path"),
        (5, 2, 900, "internal task (same root ID)"),
        (2, 6, 990, "three specialist turns in parallel"),
        (6, 5, 1080, "three correlated replies"),
        (5, 3, 1170, "resolved tool results"),
        (3, 4, 1260, "final synthesis"),
        (3, 0, 1350, "persisted answer + done"),
    ]
    for source, target, y, label in events:
        x1 = x_positions[source]
        x2 = x_positions[target]
        color = TEAL if source == 5 or target == 5 else BLUE
        _arrow(draw, (x1, y), (x2, y), color=color, width=5)
        bbox = draw.textbbox((0, 0), label, font=_font(21))
        label_width = bbox[2] - bbox[0]
        label_x = (x1 + x2) / 2 - label_width / 2
        draw.rectangle((label_x - 8, y - 36, label_x + label_width + 8, y - 6), fill=WHITE)
        draw.text((label_x, y - 35), label, font=_font(21), fill=DARK)

    image.save(path, dpi=(240, 240))


def create_grader_ablation_figure(path: Path) -> None:
    data_path = REPORT_DIR / "finance_reanalysis.json"
    data = json.loads(data_path.read_text(encoding="utf-8"))
    records = data["grader_ablation_r1_r5"]
    labels = [
        "Current",
        "No negation",
        "No hedge",
        "No conditional",
        "No token gap",
        "No plural morph.",
        "Extended inflections",
    ]
    modes = [
        ("solo_open_book", "Open-Book", "#4C78A8"),
        ("solo_two_pass", "Two-Pass", "#72B7B2"),
        ("team", "Team", "#F58518"),
        ("oracle_team", "Routed diagnostic", "#B279A2"),
    ]

    image = Image.new("RGB", (2400, 1500), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((80, 55), "Grader Sensitivity Without Additional Model Calls", font=_font(46, bold=True), fill=NAVY)
    draw.text((80, 116), "Pass rate over retained r1–r5 outputs (30 outputs per mode)", font=_font(26), fill=MUTED)

    left, top, right, bottom = 180, 250, 2320, 1240
    for tick in range(0, 101, 20):
        y = bottom - (bottom - top) * tick / 100
        draw.line((left, y, right, y), fill=MID_GRAY, width=2)
        draw.text((95, y - 13), f"{tick}%", font=_font(21), fill=MUTED)
    draw.line((left, top, left, bottom), fill=DARK, width=3)
    draw.line((left, bottom, right, bottom), fill=DARK, width=3)

    group_width = (right - left) / len(records)
    bar_width = 48
    for group_index, record in enumerate(records):
        center = left + group_width * (group_index + 0.5)
        for mode_index, (key, _, color) in enumerate(modes):
            rate = record["modes"][key]["rate"]
            x1 = center + (mode_index - 1.5) * (bar_width + 8) - bar_width / 2
            y1 = bottom - (bottom - top) * rate
            draw.rectangle((x1, y1, x1 + bar_width, bottom), fill=color)
        label = labels[group_index]
        bbox = draw.textbbox((0, 0), label, font=_font(19))
        draw.text((center - (bbox[2] - bbox[0]) / 2, bottom + 25), label, font=_font(19), fill=DARK)

    legend_x = 300
    for _, label, color in modes:
        draw.rectangle((legend_x, 1360, legend_x + 32, 1392), fill=color)
        draw.text((legend_x + 44, 1358), label, font=_font(21), fill=DARK)
        legend_x += 430
    image.save(path, dpi=(240, 240))


def create_finance_figure(path: Path) -> None:
    image = Image.new("RGB", (2400, 1450), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((80, 55), "Evidence-Grounded Financial Research Workflow", font=_font(47, bold=True), fill=NAVY)
    draw.text(
        (80, 118),
        "Deterministic evidence and state surround—not replace—model judgment",
        font=_font(26),
        fill=MUTED,
    )

    boxes = {
        "request": (90, 300, 485, 520),
        "tools": (650, 230, 1120, 520),
        "envelope": (1285, 230, 1760, 520),
        "coordinator": (1910, 300, 2310, 520),
        "skill": (650, 700, 1120, 970),
        "challenger": (1285, 700, 1760, 970),
        "ledger": (1910, 700, 2310, 970),
        "human": (1285, 1130, 1760, 1360),
    }
    _rounded_box(draw, boxes["request"], "Research Request", "symbol, market, portfolio, or event", fill=LIGHT_BLUE)
    _rounded_box(
        draw,
        boxes["tools"],
        "Deterministic Tools",
        "snapshot · screen · events · macro · risk",
        fill=LIGHT_TEAL,
    )
    _rounded_box(
        draw,
        boxes["envelope"],
        "finance.tool.v1",
        "as_of · sources · freshness · completeness · errors",
        fill=LIGHT_BLUE,
    )
    _rounded_box(draw, boxes["coordinator"], "Coordinator", "plan · verify · synthesize · disclose uncertainty", fill=LIGHT_TEAL)
    _rounded_box(draw, boxes["skill"], "Serenity Skill", "value chain · bottleneck · evidence grade · invalidation", fill=LIGHT_BLUE)
    _rounded_box(draw, boxes["challenger"], "Optional Challenger", "independent contradiction and downside review", fill=LIGHT_TEAL)
    _rounded_box(draw, boxes["ledger"], "Thesis & Alerts", "tenant scope · versioning · deduplication · audit", fill=LIGHT_BLUE)
    _rounded_box(draw, boxes["human"], "Human Decision", "approve, defer, monitor, or reject", fill="#FFF4E5", outline="#B97818")

    def edge(source: str, target: str, *, vertical: bool = False, color: str = BLUE) -> None:
        a = boxes[source]
        b = boxes[target]
        if vertical:
            start = ((a[0] + a[2]) // 2, a[3] + 10)
            end = ((b[0] + b[2]) // 2, b[1] - 10)
        else:
            start = (a[2] + 10, (a[1] + a[3]) // 2)
            end = (b[0] - 10, (b[1] + b[3]) // 2)
        _arrow(draw, start, end, color=color)

    edge("request", "tools")
    edge("tools", "envelope")
    edge("envelope", "coordinator")
    edge("tools", "skill", vertical=True, color=TEAL)
    edge("envelope", "challenger", vertical=True, color=TEAL)
    edge("coordinator", "ledger", vertical=True, color=TEAL)
    edge("challenger", "ledger")
    edge("challenger", "human", vertical=True, color="#B97818")
    draw.line((2110, 980, 2110, 1245, 1770, 1245), fill="#B97818", width=6)
    _arrow(draw, (1770, 1245), (1768, 1245), color="#B97818")

    image.save(path, dpi=(240, 240))


def create_figures() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    figures = {
        "FIGURE_1": ASSET_DIR / "fastclaw_architecture.png",
        "FIGURE_2": ASSET_DIR / "multiagent_sequence.png",
        "FIGURE_3": ASSET_DIR / "finance_workflow.png",
        "FIGURE_4": ASSET_DIR / "grader_ablation.png",
    }
    create_architecture_figure(figures["FIGURE_1"])
    create_sequence_figure(figures["FIGURE_2"])
    create_finance_figure(figures["FIGURE_3"])
    create_grader_ablation_figure(figures["FIGURE_4"])
    return figures


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill.lstrip("#"))


def _set_cell_margins(cell, top: int = 55, start: int = 70, bottom: int = 55, end: int = 70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _keep_table_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def _set_table_fixed(table) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")


def _set_table_borders(table, color: str = "B8C0C8", size: int = 5) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def _set_column_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_inches)


def _set_run_font(run, name: str = "Times New Roman", size: float | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def _configure_style_font(style, name: str, size: float, bold: bool = False, italic: bool = False) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic


def _remove_all_body_content(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _remove_header_footer(section) -> None:
    for container in (section.header, section.footer):
        element = container._element
        for child in list(element):
            element.remove(child)


def _add_bottom_border(paragraph, color: str = MID_GRAY, size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color.lstrip("#"))
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("FastClaw Undergraduate Thesis · Zheyu Wang  |  ")
    _set_run_font(prefix, size=8)
    prefix.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])
    _set_run_font(run, size=8)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.34)

    _remove_header_footer(section)
    header_p = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_p.add_run("UNDERGRADUATE THESIS  ·  AGENT RUNTIME AND FINANCIAL RESEARCH")
    _set_run_font(header_run, size=7.5)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
    _add_bottom_border(header_p, MID_GRAY, 6)

    footer_p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    _add_page_number(footer_p)

    styles = document.styles
    normal = styles["Normal"] if "Normal" in styles else styles.add_style("Normal", WD_STYLE_TYPE.PARAGRAPH)
    _configure_style_font(normal, "Times New Roman", 10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(5.0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, bold, italic in (
        ("Title", 22, True, False),
        ("Subtitle", 10.5, False, False),
        ("Heading 1", 13, True, False),
        ("Heading 2", 11.2, True, False),
        ("Heading 3", 10.5, True, True),
    ):
        if style_name in styles:
            style = styles[style_name]
        else:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        _configure_style_font(style, "Times New Roman", size, bold, italic)

    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Title"].paragraph_format.space_before = Pt(2)
    styles["Title"].paragraph_format.space_after = Pt(7)
    styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Subtitle"].paragraph_format.space_after = Pt(2)
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(8)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    styles["Heading 3"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Heading 3"].paragraph_format.space_before = Pt(7)
    styles["Heading 3"].paragraph_format.space_after = Pt(3)
    styles["Heading 3"].paragraph_format.keep_with_next = True

    for name, size, italic in (
        ("Report Abstract", 9.5, False),
        ("Report Caption", 9.0, False),
        ("Report Reference", 9.0, False),
        ("Report Code", 8.8, False),
    ):
        if name in styles:
            style = styles[name]
        else:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _configure_style_font(style, "Courier New" if name == "Report Code" else "Times New Roman", size, False, italic)
        style.paragraph_format.space_after = Pt(3.5)
        style.paragraph_format.line_spacing = 1.0
    styles["Report Abstract"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    styles["Report Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Report Caption"].paragraph_format.keep_with_next = True
    styles["Report Reference"].paragraph_format.left_indent = Inches(0.25)
    styles["Report Reference"].paragraph_format.first_line_indent = Inches(-0.25)
    styles["Report Code"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.core_properties.title = "FastClaw: Design and Evaluation of an Agent Runtime for Evidence-Grounded Financial Research"
    document.core_properties.author = "Zheyu Wang"
    document.core_properties.subject = "Undergraduate thesis on FastClaw agent runtime, financial application, and evaluation"
    document.core_properties.keywords = "agent runtime, multi-agent, tool calling, financial research, evaluation"
    update_fields = document.settings._element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        document.settings._element.append(update_fields)
    update_fields.set(qn("w:val"), "true")


INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline_text(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            _set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run)
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, "Courier New", 8.7)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_run_font(run)


def add_title_block(document: Document, metadata: dict[str, str]) -> None:
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(3)
    kicker_run = kicker.add_run("UNDERGRADUATE THESIS")
    _set_run_font(kicker_run, size=8)
    kicker_run.bold = True
    kicker_run.font.color.rgb = RGBColor.from_string(BLUE.lstrip("#"))

    title = document.add_paragraph(style="Title")
    title_run = title.add_run(metadata["TITLE"])
    _set_run_font(title_run, size=22)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(DARK.lstrip("#"))

    author = document.add_paragraph(style="Subtitle")
    author_run = author.add_run(metadata["AUTHOR"])
    _set_run_font(author_run, size=11)
    author_run.bold = True

    affiliation = document.add_paragraph(style="Subtitle")
    aff_run = affiliation.add_run(metadata["AFFILIATION"])
    _set_run_font(aff_run, size=9)
    aff_run.italic = True

    email = document.add_paragraph(style="Subtitle")
    email_run = email.add_run(f"{metadata['EMAIL']}  ·  {metadata['REPORT DATE']}")
    _set_run_font(email_run, size=9)
    email_run.italic = True
    email.paragraph_format.space_after = Pt(8)

    divider = document.add_paragraph()
    divider.paragraph_format.space_after = Pt(5)
    _add_bottom_border(divider, BLUE, 9)

    abstract = document.add_paragraph(style="Report Abstract")
    label = abstract.add_run("Abstract—")
    _set_run_font(label, size=9.1)
    label.bold = True
    add_inline_text(abstract, metadata["ABSTRACT"])

    index_terms = document.add_paragraph(style="Report Abstract")
    label = index_terms.add_run("Index Terms—")
    _set_run_font(label, size=9.1)
    label.bold = True
    add_inline_text(index_terms, metadata["INDEX TERMS"])


def _format_table_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str | None = None,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    text = text.replace("`", "").replace("**", "")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    _set_run_font(run, size=8.2)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)


def _table_alignments(separator: list[str], columns: int) -> list:
    alignments = []
    for index in range(columns):
        marker = separator[index].strip() if index < len(separator) else ""
        if marker.startswith(":") and marker.endswith(":"):
            alignments.append(WD_ALIGN_PARAGRAPH.CENTER)
        elif marker.endswith(":"):
            alignments.append(WD_ALIGN_PARAGRAPH.RIGHT)
        else:
            alignments.append(WD_ALIGN_PARAGRAPH.LEFT)
    return alignments


def _table_widths(rows: list[list[str]], columns: int) -> list[float]:
    weights = []
    for column in range(columns):
        longest = max((len(row[column]) if column < len(row) else 0) for row in rows)
        weights.append(max(8.0, min(44.0, float(longest))) ** 0.72)
    total = sum(weights) or 1.0
    widths = [BODY_WIDTH_INCHES * weight / total for weight in weights]
    minimum = 0.72 if columns >= 5 else 0.90
    widths = [max(minimum, width) for width in widths]
    scale = BODY_WIDTH_INCHES / sum(widths)
    return [width * scale for width in widths]


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    header = rows[0]
    has_separator = all(re.fullmatch(r":?-{3,}:?", item.strip()) for item in rows[1])
    body = rows[2:] if has_separator else rows[1:]
    alignments = _table_alignments(rows[1], len(header)) if has_separator else [WD_ALIGN_PARAGRAPH.LEFT] * len(header)
    table = document.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_fixed(table)
    _set_table_borders(table)
    widths = _table_widths([header, *body], len(header))
    header_row = table.rows[0]
    _set_repeat_table_header(header_row)
    _keep_table_row_together(header_row)
    for index, value in enumerate(header):
        _set_column_width(header_row.cells[index], widths[index])
        _set_cell_shading(header_row.cells[index], NAVY)
        _format_table_text(header_row.cells[index], value, bold=True, color=WHITE, alignment=alignments[index])
    for row_values in body:
        row = table.add_row()
        _keep_table_row_together(row)
        for index, value in enumerate(row_values):
            _set_column_width(row.cells[index], widths[index])
            if len(table.rows) % 2 == 1:
                _set_cell_shading(row.cells[index], LIGHT_GRAY)
            _format_table_text(row.cells[index], value, alignment=alignments[index])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def _roman(number: int) -> str:
    values = (
        (1000, "M"),
        (900, "CM"),
        (500, "D"),
        (400, "CD"),
        (100, "C"),
        (90, "XC"),
        (50, "L"),
        (40, "XL"),
        (10, "X"),
        (9, "IX"),
        (5, "V"),
        (4, "IV"),
        (1, "I"),
    )
    result = ""
    remaining = number
    for value, numeral in values:
        while remaining >= value:
            result += numeral
            remaining -= value
    return result


def add_table_caption(document: Document, number: int, text: str) -> None:
    caption = document.add_paragraph(style="Report Caption")
    caption.paragraph_format.keep_with_next = True
    run = caption.add_run(f"Table {_roman(number)}. {text}")
    _set_run_font(run, size=8.5)
    run.bold = True


def add_figure(document: Document, image_path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(BODY_WIDTH_INCHES))
    caption_p = document.add_paragraph(style="Report Caption")
    caption_run = caption_p.add_run(caption)
    _set_run_font(caption_run, size=8.5)
    caption_run.italic = True


def add_algorithm(document: Document, caption_text: str, steps: list[str]) -> None:
    caption = document.add_paragraph(style="Report Caption")
    caption_run = caption.add_run(f"Algorithm 1. {caption_text}")
    _set_run_font(caption_run, size=8.5)
    caption_run.bold = True
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_fixed(table)
    _set_table_borders(table)
    cell = table.cell(0, 0)
    _keep_table_row_together(table.rows[0])
    _set_column_width(cell, BODY_WIDTH_INCHES)
    _set_cell_shading(cell, "F7F9FB")
    _set_cell_margins(cell, 90, 120, 90, 120)
    for index, step in enumerate(steps):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1.5)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(step)
        _set_run_font(run, "Courier New", 8.2)


def add_body_paragraph(document: Document, text: str) -> None:
    if text.startswith("[") and re.match(r"^\[\d+\]", text):
        paragraph = document.add_paragraph(style="Report Reference")
    elif text.startswith("`") and text.endswith("`"):
        paragraph = document.add_paragraph(style="Report Code")
    else:
        paragraph = document.add_paragraph(style="Normal")
    add_inline_text(paragraph, text)


def add_list_item(document: Document, text: str, *, numbered: bool = False) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_after = Pt(2.2)
    marker = ""
    body = text
    if numbered:
        match = re.match(r"^(\d+)\.\s+(.*)$", text)
        if match:
            marker = f"{match.group(1)}. "
            body = match.group(2)
    else:
        marker = "• "
    marker_run = paragraph.add_run(marker)
    _set_run_font(marker_run)
    marker_run.bold = numbered
    add_inline_text(paragraph, body)


def _add_field(paragraph, instruction: str, placeholder: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field_code = OxmlElement("w:instrText")
    field_code.set(qn("xml:space"), "preserve")
    field_code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, field_code, separate, text, end])
    _set_run_font(run, size=10.5)


def _front_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    _set_run_font(run, size=14)
    run.bold = True


def _front_entry(document: Document, label: str, page: int | str | None, *, indent: float = 0.0) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(indent)
    paragraph.paragraph_format.space_after = Pt(1.8)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(BODY_WIDTH_INCHES - 0.08),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )
    add_inline_text(paragraph, label)
    page_run = paragraph.add_run(f"\t{page if page is not None else '—'}")
    _set_run_font(page_run, size=9.5)


def add_front_matter(
    document: Document,
    abbreviations: list[tuple[str, str]],
    lines: list[str],
    figure_captions: dict[str, str],
    page_map: dict[str, int],
) -> None:
    document.add_page_break()
    _front_heading(document, "TABLE OF CONTENTS")
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("# "):
            heading = stripped[2:]
            _front_entry(document, heading, page_map.get(f"heading::{heading}"))

    document.add_page_break()
    _front_heading(document, "LIST OF FIGURES")
    for key in sorted(figure_captions):
        _front_entry(document, figure_captions[key], page_map.get(f"figure::{key}"), indent=0.15)

    _front_heading(document, "LIST OF TABLES")
    table_number = 1
    for line in lines:
        if line.strip().startswith("TABLE_CAPTION:"):
            _front_entry(
                document,
                f"Table {_roman(table_number)}. {line.split(':', 1)[1].strip()}",
                page_map.get(f"table::{table_number}"),
                indent=0.15,
            )
            table_number += 1

    document.add_page_break()
    _front_heading(document, "LIST OF ABBREVIATIONS")
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_fixed(table)
    _set_table_borders(table)
    widths = [1.45, BODY_WIDTH_INCHES - 1.45]
    for index, value in enumerate(("Abbreviation", "Meaning")):
        _set_column_width(table.rows[0].cells[index], widths[index])
        _set_cell_shading(table.rows[0].cells[index], NAVY)
        _format_table_text(table.rows[0].cells[index], value, bold=True, color=WHITE)
    _set_repeat_table_header(table.rows[0])
    for abbreviation, meaning in abbreviations:
        row = table.add_row()
        _keep_table_row_together(row)
        for index, value in enumerate((abbreviation, meaning)):
            _set_column_width(row.cells[index], widths[index])
            _format_table_text(row.cells[index], value)
    document.add_page_break()


def parse_source() -> tuple[dict[str, str], list[tuple[str, str]], list[str]]:
    lines = SOURCE_PATH.read_text(encoding="utf-8").splitlines()
    metadata: dict[str, str] = {}
    abbreviations: list[tuple[str, str]] = []
    body_start = 0
    for index, line in enumerate(lines):
        if line.startswith("# "):
            body_start = index
            break
        match = re.match(r"^(TITLE|AUTHOR|AFFILIATION|EMAIL|REPORT DATE|ABSTRACT|INDEX TERMS):\s*(.*)$", line)
        if match:
            metadata[match.group(1)] = match.group(2).strip()
            continue
        abbreviation = re.match(r"^ABBREVIATION:\s*([^|]+)\|\s*(.*)$", line)
        if abbreviation:
            abbreviations.append((abbreviation.group(1).strip(), abbreviation.group(2).strip()))
    required = {"TITLE", "AUTHOR", "AFFILIATION", "EMAIL", "REPORT DATE", "ABSTRACT", "INDEX TERMS"}
    missing = required - metadata.keys()
    if missing:
        raise ValueError(f"Missing report metadata: {sorted(missing)}")
    return metadata, abbreviations, lines[body_start:]


def build_document() -> None:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Reference draft not found: {TEMPLATE_PATH}")
    figures = create_figures()
    working_copy = REPORT_DIR / ".working-template.docx"
    shutil.copyfile(TEMPLATE_PATH, working_copy)
    document = Document(working_copy)
    _remove_all_body_content(document)
    configure_document(document)

    metadata, abbreviations, lines = parse_source()
    add_title_block(document, metadata)

    figure_captions = {
        "FIGURE_1": "Fig. 1. FastClaw layered architecture and the principal control boundaries.",
        "FIGURE_2": "Fig. 2. One coordinator turn through TaskQueue and MessageBus with parallel specialist delegation.",
        "FIGURE_3": "Fig. 3. Financial research workflow separating evidence, model judgment, and human action.",
        "FIGURE_4": "Fig. 4. Pass-rate sensitivity to six lexical grader design choices over retained r1–r5 outputs.",
    }
    page_map = json.loads(PAGE_MAP_PATH.read_text(encoding="utf-8")) if PAGE_MAP_PATH.exists() else {}
    add_front_matter(document, abbreviations, lines, figure_captions, page_map)

    index = 0
    table_number = 1
    chapter_number = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Heading 1")
            if chapter_number > 0:
                paragraph.paragraph_format.page_break_before = True
            run = paragraph.add_run(line[2:].upper())
            _set_run_font(run, size=13)
            run.bold = True
            chapter_number += 1
            index += 1
            continue
        if line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 2")
            run = paragraph.add_run(line[3:])
            _set_run_font(run, size=11.2)
            run.bold = True
            index += 1
            continue
        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 3")
            run = paragraph.add_run(line[4:])
            _set_run_font(run, size=10.5)
            run.bold = True
            run.italic = True
            index += 1
            continue
        if line.startswith("TABLE_CAPTION:"):
            add_table_caption(document, table_number, line.split(":", 1)[1].strip())
            table_number += 1
            index += 1
            continue
        placeholder = re.fullmatch(r"\[\[([A-Z0-9_]+)\]\]", line)
        if placeholder:
            key = placeholder.group(1)
            if key in figures:
                add_figure(document, figures[key], figure_captions[key])
            elif key == "ALGORITHM_1":
                pass
            elif key == "PAGE_BREAK":
                document.add_page_break()
            else:
                raise ValueError(f"Unknown placeholder: {key}")
            index += 1
            continue
        if line.startswith("ALGORITHM_CAPTION:"):
            caption = line.split(":", 1)[1].strip()
            steps: list[str] = []
            index += 1
            while index < len(lines):
                candidate = lines[index].strip()
                if not candidate:
                    index += 1
                    continue
                if not candidate.startswith("ALGORITHM_STEP:"):
                    break
                steps.append(candidate.split(":", 1)[1].strip())
                index += 1
            add_algorithm(document, caption, steps)
            continue
        if line.startswith("|"):
            table_rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append([item.strip() for item in lines[index].strip().strip("|").split("|")])
                index += 1
            add_markdown_table(document, table_rows)
            continue
        if re.match(r"^\d+\.\s+", line):
            add_list_item(document, line, numbered=True)
            index += 1
            continue
        if line.startswith("- "):
            add_list_item(document, line[2:], numbered=False)
            index += 1
            continue
        add_body_paragraph(document, line)
        index += 1

    document.save(OUTPUT_PATH)
    working_copy.unlink(missing_ok=True)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
